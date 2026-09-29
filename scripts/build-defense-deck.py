#!/usr/bin/env python3
"""Build the thesis defense presentation from `defense_content.py`.

Usage:
    uv run python scripts/build-defense-deck.py
    uv run python scripts/build-defense-deck.py --check

Every shape is positioned absolutely, in inches, against the token table
below. There is no slide master beyond a blank layout, so what this script
writes is exactly what shows on screen.

The figures under `documents/deck-figures/` were rasterised once from the
vector figures the notebooks produce. Refresh one with:
    pdftoppm -r 200 -png -singlefile thesis/figures/<name>.pdf documents/deck-figures/<name>

`--check` catches overflow and collisions arithmetically, but it cannot see
the real font metrics. To actually look at the slides, render them:
    soffice --headless --convert-to pdf --outdir /tmp/deck documents/ensights-thesis-defense.pptx
    pdftoppm -png -r 80 /tmp/deck/ensights-thesis-defense.pdf /tmp/deck/slide
Do this before presenting. It is how the title-slide wrap and the divider
numeral collision were found, and neither was visible in the file itself.
"""

import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu
from pptx.util import Inches
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).parent))

from defense_content import SECTIONS  # noqa: E402
from defense_content import SLIDES  # noqa: E402

PROJECT_ROOT = Path(__file__).parent.parent
OUTPUT = PROJECT_ROOT / "documents" / "ensights-thesis-defense.pptx"

# Palette.
GREEN = RGBColor(0x2C, 0x55, 0x30)
GREEN_LIGHT = RGBColor(0xEE, 0xF3, 0xEC)
GREEN_MID = RGBColor(0xD7, 0xE5, 0xD2)
AMBER = RGBColor(0xE8, 0xA3, 0x3D)
RED = RGBColor(0xB0, 0x41, 0x3E)
INK = RGBColor(0x24, 0x1F, 0x1B)
MUTED = RGBColor(0x6B, 0x62, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FRAME = RGBColor(0xD8, 0xDE, 0xEC)

# Typography.
HEADING_FONT = "Cambria"
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

# Geometry, in inches.
SLIDE_WIDTH = 13.333
SLIDE_HEIGHT = 7.5
MARGIN = 0.60
CONTENT_WIDTH = 12.13
COLUMN_WIDTH = 5.85
RIGHT_COLUMN = 6.88
CONTENT_TOP = 1.62
CONTENT_BOTTOM = 6.86
FOOTER_TOP = 7.08
BLOCK_GAP = 0.16

# Average glyph width as a fraction of the point size, used to guess how many
# characters fit on a line. The heading figure is calibrated against a real
# LibreOffice render: the 59-character title slide heading wraps at 31pt and
# the 52-character heading on slide 9 does not at 30pt, which brackets it
# between 0.478 and 0.560.
GLYPH_WIDTH = {HEADING_FONT: 0.53, BODY_FONT: 0.48, MONO_FONT: 0.55}
LINE_SPACING = 1.22

# Register the earlier draft of this deck fell into: translationese, noun
# stacks and symbols standing in for words. None of it is how anyone speaks,
# so the check fails the build if any of it comes back.
STIFF_PHRASES = (
    "⇒",
    "τα δύο αλγόριθμοι",
    "Ό,τι επιτρέπει",
    "αναπαραγώγιμ",
    "ημιτελ",
    "εκ των προτέρων",
    "διαμερίζ",
    "τεκμήριο",
    "κλειστό-προς-τα-κάτω",
    "Ρητή διάκριση",
    "απόλυτα κατώφλια",
    "προσάρτηση",
    "δεσμευμένες συγγραφές",
)

warnings: list[str] = []


def wrapped_height(text, width, size, font=BODY_FONT):
    """Estimate the rendered height of `text` wrapped into `width` inches."""
    characters_per_line = max(1, (width * 72) / (size * GLYPH_WIDTH[font]))
    lines = max(1, math.ceil(len(text) / characters_per_line))
    return lines * size * LINE_SPACING / 72


def add_textbox(slide, x, y, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    return frame


def write(
    frame,
    text,
    size,
    colour,
    font=BODY_FONT,
    bold=False,
    align=PP_ALIGN.LEFT,
    space_after=0,
    first=False,
):
    """Append a paragraph to `frame`, or fill its first empty one."""
    paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
    paragraph.alignment = align
    paragraph.space_after = Pt(space_after)
    paragraph.line_spacing = LINE_SPACING
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = colour
    return paragraph


def simple_text(slide, x, y, width, height, text, size, colour, **kwargs):
    frame = add_textbox(slide, x, y, width, height)
    write(frame, text, size, colour, first=True, **kwargs)
    return frame


def must_fit(label, text, width, height, size, font):
    """Warn when a fixed-height box is too short for the text it will hold.

    Headings sit in hand-placed boxes with something directly underneath, so
    an extra wrapped line does not push the next block down, it prints on top
    of it. That is invisible in the file and only shows up once rendered.
    """
    needed = wrapped_height(text, width, size, font)
    if needed > height + 0.02:
        warnings.append(f"{label}: needs {needed:.2f}in but the box is {height:.2f}in")


def add_rectangle(
    slide, x, y, width, height, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None
):
    box = slide.shapes.add_shape(
        shape, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    box.shadow.inherit = False
    if fill is None:
        box.fill.background()
    else:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = line
        box.line.width = Pt(0.75)
    box.text_frame.word_wrap = True
    return box


# --------------------------------------------------------------------------
# Blocks. Each renderer draws itself at (x, y) across `width` and returns the
# height it consumed, so a column can stack blocks without knowing their kind.
# --------------------------------------------------------------------------


def render_paragraphs(slide, block, x, y, width):
    size = block.get("size", 15)
    frame = add_textbox(slide, x, y, width, 0.4)
    height = 0.0
    for index, item in enumerate(block["items"]):
        text = item["text"] if isinstance(item, dict) else item
        bold = isinstance(item, dict) and item.get("bold", False)
        write(frame, text, size, INK, bold=bold, space_after=9, first=index == 0)
        height += wrapped_height(text, width, size) + 9 / 72
    return height


def render_bullets(slide, block, x, y, width):
    size = block.get("size", 15)
    height = 0.0
    if block.get("label"):
        simple_text(slide, x, y, width, 0.28, block["label"], 12, AMBER, bold=True)
        height += 0.40
    frame = add_textbox(slide, x, y + height, width, 0.4)
    for index, item in enumerate(block["items"]):
        text = f"—  {item}"
        write(frame, text, size, INK, space_after=10, first=index == 0)
        height += wrapped_height(text, width, size) + 10 / 72
    return height


def render_cards(slide, block, x, y, width):
    # Fill, heading colour, body colour, outline. The alert card is white, so
    # it needs an outline or it disappears into the slide.
    tones = {
        "light": (GREEN_LIGHT, GREEN, INK, None),
        "mid": (GREEN_MID, GREEN, INK, None),
        "solid": (GREEN, WHITE, GREEN_MID, None),
        "alert": (WHITE, RED, RED, RED),
    }
    columns = block.get("columns", 1)
    fixed = block.get("height")
    gap = 0.20
    card_width = (width - gap * (columns - 1)) / columns

    heights = []
    for card in block["items"]:
        body = wrapped_height(card["text"], card_width - 0.50, 14)
        heights.append(fixed or max(1.15, 0.66 + body + 0.24))
    row_height = max(heights)

    for index, card in enumerate(block["items"]):
        column = index % columns
        row = index // columns
        card_x = x + column * (card_width + gap)
        card_y = y + row * (row_height + gap)
        fill, heading_colour, text_colour, outline = tones[card.get("tone", "light")]
        add_rectangle(slide, card_x, card_y, card_width, row_height, fill, line=outline)
        simple_text(
            slide,
            card_x + 0.25,
            card_y + 0.18,
            card_width - 0.50,
            0.30,
            card["heading"],
            13.5,
            heading_colour,
            font=HEADING_FONT,
            bold=True,
        )
        simple_text(
            slide,
            card_x + 0.25,
            card_y + 0.55,
            card_width - 0.50,
            row_height - 0.76,
            card["text"],
            14,
            text_colour,
        )

    rows = math.ceil(len(block["items"]) / columns)
    return rows * row_height + (rows - 1) * gap


def render_chips(slide, block, x, y, width):
    height = 0.0
    if block.get("label"):
        simple_text(slide, x, y, width, 0.28, block["label"], 12, AMBER, bold=True)
        height += 0.38
    gap = 0.14
    chip_width = (width - gap) / 2
    for index, item in enumerate(block["items"]):
        column = index % 2
        row = index // 2
        chip_x = x + column * (chip_width + gap)
        chip_y = y + height + row * (0.42 + 0.10)
        add_rectangle(slide, chip_x, chip_y, chip_width, 0.42, GREEN_LIGHT)
        simple_text(
            slide,
            chip_x,
            chip_y + 0.10,
            chip_width,
            0.26,
            item,
            12,
            GREEN,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
    rows = math.ceil(len(block["items"]) / 2)
    return height + rows * 0.42 + (rows - 1) * 0.10


def render_bars(slide, block, x, y, width):
    height = 0.0
    for label, description in block["items"]:
        add_rectangle(slide, x, y + height, width, 0.62, GREEN_LIGHT)
        add_rectangle(
            slide, x, y + height, 0.09, 0.62, GREEN, shape=MSO_SHAPE.RECTANGLE
        )
        simple_text(
            slide,
            x + 0.28,
            y + height + 0.09,
            2.0,
            0.26,
            label,
            13,
            GREEN,
            font=HEADING_FONT,
            bold=True,
        )
        simple_text(
            slide,
            x + 2.30,
            y + height + 0.11,
            width - 2.55,
            0.26,
            description,
            12,
            MUTED,
        )
        height += 0.62 + 0.12
    return height - 0.12


def render_table(slide, block, x, y, width):
    height = 0.0
    if block.get("label"):
        simple_text(slide, x, y, width, 0.28, block["label"], 11.5, AMBER, bold=True)
        height += 0.36

    headers = block["headers"]
    widths = block.get("widths") or [1 / len(headers)] * len(headers)
    row_height = 0.40
    header_height = 0.38

    add_rectangle(
        slide, x, y + height, width, header_height, GREEN_MID, shape=MSO_SHAPE.RECTANGLE
    )
    offset = 0.0
    for header, share in zip(headers, widths, strict=True):
        simple_text(
            slide,
            x + offset + 0.10,
            y + height + 0.09,
            width * share - 0.20,
            0.26,
            header,
            12,
            GREEN,
            bold=True,
        )
        offset += width * share
    height += header_height

    for index, row in enumerate(block["rows"]):
        if index % 2 == 1:
            add_rectangle(
                slide,
                x,
                y + height,
                width,
                row_height,
                GREEN_LIGHT,
                shape=MSO_SHAPE.RECTANGLE,
            )
        offset = 0.0
        for cell, share in zip(row, widths, strict=True):
            simple_text(
                slide,
                x + offset + 0.10,
                y + height + 0.10,
                width * share - 0.20,
                0.26,
                str(cell),
                12.5,
                INK,
            )
            offset += width * share
        height += row_height
    return height


def render_image(slide, block, x, y, width):
    path = PROJECT_ROOT / block["path"]
    if not path.exists():
        raise FileNotFoundError(path)
    picture = slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), height=Inches(block["height"])
    )
    actual_width = picture.width / 914400
    if actual_width > width:
        # These charts are wider than they are tall, so the column, not the
        # requested height, is usually what limits them.
        picture.height = Inches(width / actual_width * picture.height / 914400)
        picture.width = Inches(width)
        actual_width = width
    picture.left = Inches(x + (width - actual_width) / 2)

    # `box` reserves a fixed slot so two images of different proportions,
    # side by side, still line their captions up.
    drawn_height = picture.height / 914400
    box = max(block.get("box", 0), drawn_height)
    picture.top = Inches(y + (box - drawn_height) / 2)

    add_rectangle(
        slide,
        picture.left / 914400,
        picture.top / 914400,
        picture.width / 914400,
        drawn_height,
        None,
        shape=MSO_SHAPE.RECTANGLE,
        line=FRAME,
    )
    height = box
    if block.get("caption"):
        simple_text(
            slide, x, y + height + 0.10, width, 0.32, block["caption"], 11, MUTED
        )
        height += 0.10 + wrapped_height(block["caption"], width, 11)
    return height


def render_statistics(slide, block, x, y, width):
    columns = block.get("columns", 3)
    gap = 0.16
    tile_width = (width - gap * (columns - 1)) / columns
    tile_height = 1.06
    for index, (value, label) in enumerate(block["items"]):
        tile_x = x + (index % columns) * (tile_width + gap)
        tile_y = y + (index // columns) * (tile_height + gap)
        add_rectangle(slide, tile_x, tile_y, tile_width, tile_height, GREEN_LIGHT)
        simple_text(
            slide,
            tile_x,
            tile_y + 0.10,
            tile_width,
            0.42,
            value,
            24,
            GREEN,
            font=HEADING_FONT,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        simple_text(
            slide,
            tile_x + 0.06,
            tile_y + 0.60,
            tile_width - 0.12,
            0.36,
            label,
            11,
            MUTED,
            align=PP_ALIGN.CENTER,
        )
    rows = math.ceil(len(block["items"]) / columns)
    return rows * tile_height + (rows - 1) * gap


def render_code(slide, block, x, y, width):
    simple_text(
        slide,
        x,
        y,
        width,
        0.28,
        block["heading"],
        11.5,
        AMBER,
        font=MONO_FONT,
        bold=True,
    )
    height = 0.38
    panel_height = len(block["lines"]) * 0.225 + 0.32
    add_rectangle(
        slide,
        x,
        y + height,
        width,
        panel_height,
        GREEN_LIGHT,
        shape=MSO_SHAPE.RECTANGLE,
    )
    frame = add_textbox(
        slide, x + 0.16, y + height + 0.14, width - 0.32, panel_height - 0.28
    )
    for index, line in enumerate(block["lines"]):
        write(frame, line, 11.5, INK, font=MONO_FONT, first=index == 0)
    height += panel_height
    if block.get("caption"):
        simple_text(
            slide, x, y + height + 0.10, width, 0.40, block["caption"], 10.5, MUTED
        )
        height += 0.10 + wrapped_height(block["caption"], width, 10.5)
    return height


def render_note(slide, block, x, y, width):
    text_height = wrapped_height(block["text"], width - 0.50, 12.5)
    height = max(0.64, text_height + 0.36)
    add_rectangle(
        slide, x, y, width, height, None, shape=MSO_SHAPE.RECTANGLE, line=GREEN_MID
    )
    add_rectangle(slide, x, y, 0.07, height, AMBER, shape=MSO_SHAPE.RECTANGLE)
    simple_text(
        slide,
        x + 0.25,
        y + 0.16,
        width - 0.45,
        height - 0.32,
        block["text"],
        12.5,
        MUTED,
    )
    return height


def render_numbered(slide, block, x, y, width):
    columns = 2
    gap = 0.30
    item_width = (width - gap) / columns
    row_height = 0.92
    for index, item in enumerate(block["items"]):
        column = index % columns
        row = index // columns
        item_x = x + column * (item_width + gap)
        item_y = y + row * row_height
        add_rectangle(slide, item_x, item_y, 0.54, 0.54, GREEN, shape=MSO_SHAPE.OVAL)
        simple_text(
            slide,
            item_x,
            item_y + 0.13,
            0.54,
            0.30,
            str(index + 1),
            16,
            WHITE,
            font=HEADING_FONT,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        simple_text(
            slide, item_x + 0.74, item_y + 0.14, item_width - 0.80, 0.44, item, 16, INK
        )
    rows = math.ceil(len(block["items"]) / columns)
    return rows * row_height


RENDERERS = {
    "paragraphs": render_paragraphs,
    "bullets": render_bullets,
    "cards": render_cards,
    "chips": render_chips,
    "bars": render_bars,
    "table": render_table,
    "image": render_image,
    "statistics": render_statistics,
    "code": render_code,
    "note": render_note,
    "numbered": render_numbered,
}


def render_column(slide, blocks, x, y, width, label):
    """Stack `blocks` down a column and warn if they run past the content area."""
    cursor = y
    for block in blocks:
        cursor += RENDERERS[block["type"]](slide, block, x, cursor, width) + BLOCK_GAP
    cursor -= BLOCK_GAP
    if cursor > CONTENT_BOTTOM:
        warnings.append(
            f"{label}: content reaches {cursor:.2f}in, past {CONTENT_BOTTOM}in"
        )
    return cursor


def nudge_down(slide, first, distance):
    """Move every shape added from index `first` onwards down by `distance`."""
    if distance <= 0:
        return
    offset = Emu(int(distance * 914400))
    for shape in list(slide.shapes)[first:]:
        shape.top += offset


# --------------------------------------------------------------------------
# Slide layouts.
# --------------------------------------------------------------------------


def add_chrome(slide, spec, number):
    if spec.get("title"):
        must_fit(
            f"slide {number} title",
            spec["title"],
            CONTENT_WIDTH,
            0.70,
            30,
            HEADING_FONT,
        )
    if spec.get("eyebrow"):
        simple_text(
            slide,
            MARGIN,
            0.50,
            CONTENT_WIDTH,
            0.30,
            spec["eyebrow"],
            12,
            AMBER,
            bold=True,
        )
    if spec.get("title"):
        simple_text(
            slide,
            MARGIN,
            0.82,
            CONTENT_WIDTH,
            0.70,
            spec["title"],
            30,
            GREEN,
            font=HEADING_FONT,
            bold=True,
        )
    simple_text(slide, MARGIN, FOOTER_TOP, 4.0, 0.30, "ENSIGHTS", 9, MUTED)
    simple_text(
        slide,
        SLIDE_WIDTH - MARGIN - 0.60,
        FOOTER_TOP,
        0.60,
        0.30,
        number,
        9,
        MUTED,
        align=PP_ALIGN.RIGHT,
    )


def add_footnote(slide, text, label):
    height = wrapped_height(text, CONTENT_WIDTH, 11)
    top = CONTENT_BOTTOM + 0.06 - height
    simple_text(slide, MARGIN, top, CONTENT_WIDTH, height + 0.10, text, 11, MUTED)
    return top


def layout_title(slide, spec):
    must_fit(
        "title slide heading", spec["title"], CONTENT_WIDTH, 1.35, 31, HEADING_FONT
    )
    add_rectangle(slide, 0, 0, SLIDE_WIDTH, 3.05, GREEN, shape=MSO_SHAPE.RECTANGLE)
    simple_text(
        slide, MARGIN, 0.55, CONTENT_WIDTH, 0.30, spec["kicker"], 12, AMBER, bold=True
    )
    simple_text(
        slide,
        MARGIN,
        0.95,
        CONTENT_WIDTH,
        1.35,
        spec["title"],
        31,
        WHITE,
        font=HEADING_FONT,
        bold=True,
    )
    simple_text(
        slide, MARGIN, 2.45, CONTENT_WIDTH, 0.34, spec["subtitle"], 14, GREEN_MID
    )

    simple_text(
        slide,
        MARGIN,
        3.35,
        CONTENT_WIDTH,
        0.44,
        spec["author"],
        22,
        GREEN,
        font=HEADING_FONT,
        bold=True,
    )
    frame = add_textbox(slide, MARGIN, 3.92, CONTENT_WIDTH, 0.90)
    for index, line in enumerate(spec["programme"]):
        write(frame, line, 12, MUTED, space_after=4, first=index == 0)

    add_rectangle(slide, MARGIN, 4.90, 7.20, 1.70, GREEN_LIGHT)
    simple_text(
        slide,
        MARGIN + 0.30,
        5.08,
        6.60,
        0.28,
        "ΤΡΙΜΕΛΗΣ ΕΞΕΤΑΣΤΙΚΗ ΕΠΙΤΡΟΠΗ",
        11,
        AMBER,
        bold=True,
    )
    for index, (role, name) in enumerate(spec["committee"]):
        row_y = 5.46 + index * 0.34
        simple_text(slide, MARGIN + 0.30, row_y, 1.20, 0.26, role, 12, GREEN, bold=True)
        simple_text(slide, MARGIN + 1.55, row_y, 5.30, 0.26, name, 12, INK)

    simple_text(slide, MARGIN, 6.85, CONTENT_WIDTH, 0.30, spec["date"], 12, MUTED)


def layout_contents(slide, spec, number):
    add_chrome(slide, spec, number)
    gap = 0.24
    card_width = (CONTENT_WIDTH - gap * 3) / 4
    for index, (marker, name, blurb) in enumerate(SECTIONS):
        card_x = MARGIN + index * (card_width + gap)
        add_rectangle(slide, card_x, 2.10, card_width, 3.10, GREEN_LIGHT)
        add_rectangle(
            slide, card_x, 2.10, card_width, 0.09, GREEN, shape=MSO_SHAPE.RECTANGLE
        )
        simple_text(
            slide,
            card_x + 0.28,
            2.45,
            card_width - 0.56,
            0.70,
            marker,
            40,
            GREEN_MID,
            font=HEADING_FONT,
            bold=True,
        )
        simple_text(
            slide,
            card_x + 0.28,
            3.25,
            card_width - 0.56,
            0.50,
            name,
            17,
            GREEN,
            font=HEADING_FONT,
            bold=True,
        )
        simple_text(
            slide, card_x + 0.28, 3.85, card_width - 0.56, 1.10, blurb, 12, MUTED
        )


def layout_divider(slide, spec, number):
    marker, name, blurb = SECTIONS[spec["section"]]
    must_fit(f"divider {number} name", name, CONTENT_WIDTH, 0.80, 44, HEADING_FONT)
    add_rectangle(
        slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, GREEN, shape=MSO_SHAPE.RECTANGLE
    )
    simple_text(
        slide,
        MARGIN,
        1.75,
        CONTENT_WIDTH,
        1.80,
        marker,
        96,
        RGBColor(0x3E, 0x6E, 0x43),
        font=HEADING_FONT,
        bold=True,
    )
    simple_text(
        slide, MARGIN, 3.80, CONTENT_WIDTH, 0.32, "ΕΝΟΤΗΤΑ", 13, AMBER, bold=True
    )
    simple_text(
        slide,
        MARGIN,
        4.16,
        CONTENT_WIDTH,
        0.80,
        name,
        44,
        WHITE,
        font=HEADING_FONT,
        bold=True,
    )
    simple_text(slide, MARGIN, 5.08, CONTENT_WIDTH, 0.40, blurb, 15, GREEN_MID)
    simple_text(
        slide,
        MARGIN,
        FOOTER_TOP,
        6.0,
        0.30,
        "ENSIGHTS · Διπλωματική εργασία",
        9,
        GREEN_MID,
    )
    simple_text(
        slide,
        SLIDE_WIDTH - MARGIN - 0.60,
        FOOTER_TOP,
        0.60,
        0.30,
        number,
        9,
        GREEN_MID,
        align=PP_ALIGN.RIGHT,
    )


def layout_closing(slide, spec, number):
    add_rectangle(
        slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, GREEN, shape=MSO_SHAPE.RECTANGLE
    )
    add_rectangle(slide, MARGIN, 2.20, 0.10, 1.60, AMBER, shape=MSO_SHAPE.RECTANGLE)
    frame = add_textbox(slide, MARGIN + 0.45, 2.25, CONTENT_WIDTH - 0.45, 1.60)
    for index, line in enumerate(spec["message"]):
        write(
            frame, line, 27, WHITE, font=HEADING_FONT, space_after=6, first=index == 0
        )
    simple_text(
        slide, MARGIN + 0.45, 4.40, CONTENT_WIDTH, 0.50, spec["thanks"], 20, GREEN_MID
    )
    simple_text(
        slide,
        MARGIN + 0.45,
        5.15,
        CONTENT_WIDTH,
        0.60,
        spec["prompt"],
        34,
        AMBER,
        font=HEADING_FONT,
        bold=True,
    )
    simple_text(
        slide,
        MARGIN,
        FOOTER_TOP,
        6.0,
        0.30,
        f"ENSIGHTS · {spec.get('author', 'Χρήστος Μουρτογιάννης')}",
        9,
        GREEN_MID,
    )
    simple_text(
        slide,
        SLIDE_WIDTH - MARGIN - 0.60,
        FOOTER_TOP,
        0.60,
        0.30,
        number,
        9,
        GREEN_MID,
        align=PP_ALIGN.RIGHT,
    )


def layout_content(slide, spec, number, label):
    add_chrome(slide, spec, number)
    limit = CONTENT_BOTTOM
    if spec.get("footnote"):
        limit = add_footnote(slide, spec["footnote"], label) - 0.14

    first_content_shape = len(slide.shapes)
    bottom = CONTENT_TOP
    if spec["layout"] == "full":
        bottom = render_column(
            slide, spec["body"], MARGIN, CONTENT_TOP, CONTENT_WIDTH, label
        )
    else:
        left = render_column(
            slide, spec["left"], MARGIN, CONTENT_TOP, COLUMN_WIDTH, f"{label} left"
        )
        right = render_column(
            slide,
            spec["right"],
            RIGHT_COLUMN,
            CONTENT_TOP,
            COLUMN_WIDTH,
            f"{label} right",
        )
        bottom = max(left, right)
        if spec.get("below"):
            bottom = render_column(
                slide,
                spec["below"],
                MARGIN,
                bottom + 0.22,
                CONTENT_WIDTH,
                f"{label} below",
            )

    # Sparse slides look accidental when everything hugs the title. Centre the
    # content in the space it has, but only a little, so it never floats away
    # from the heading it belongs to.
    slack = min((limit - bottom) / 2, 0.85)
    nudge_down(slide, first_content_shape, slack)
    bottom += max(0.0, slack)

    if bottom > limit:
        warnings.append(
            f"{label}: content ends at {bottom:.2f}in but the limit is {limit:.2f}in"
        )


def notes_text(spec):
    notes = spec["notes"]
    parts = [notes["plain"]]
    if notes.get("maths"):
        parts.append(f"ΜΑΘΗΜΑΤΙΚΑ: {notes['maths']}")
    if notes.get("question"):
        parts.append(f"ΠΙΘΑΝΗ ΕΡΩΤΗΣΗ: {notes['question']}")
    return "\n\n".join(parts)


def build():
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH)
    presentation.slide_height = Inches(SLIDE_HEIGHT)
    blank = presentation.slide_layouts[6]

    for number, spec in enumerate(SLIDES, start=1):
        slide = presentation.slides.add_slide(blank)
        label = f"slide {number}"
        layout = spec["layout"]
        if layout == "title":
            layout_title(slide, spec)
        elif layout == "contents":
            layout_contents(slide, spec, str(number))
        elif layout == "divider":
            layout_divider(slide, spec, str(number))
        elif layout == "closing":
            layout_closing(slide, spec, str(number))
        else:
            layout_content(slide, spec, str(number), label)
        slide.notes_slide.notes_text_frame.text = notes_text(spec)

    presentation.save(OUTPUT)
    return presentation


def check():
    """Re-open the built deck and assert what cannot be seen without a renderer."""
    presentation = Presentation(OUTPUT)
    problems = list(warnings)

    if len(presentation.slides) != len(SLIDES):
        problems.append(
            f"expected {len(SLIDES)} slides, found {len(presentation.slides)}"
        )

    for index, slide in enumerate(presentation.slides, start=1):
        notes = slide.notes_slide.notes_text_frame.text
        if not notes.strip():
            problems.append(f"slide {index}: empty speaker notes")
        spoken = (
            " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
            + " "
            + notes
        )
        for phrase in STIFF_PHRASES:
            if phrase in spoken:
                problems.append(f"slide {index}: stiff phrase {phrase!r}")
        for shape in slide.shapes:
            right = (shape.left + shape.width) / 914400
            bottom = (shape.top + shape.height) / 914400
            if right > SLIDE_WIDTH + 0.02 or bottom > SLIDE_HEIGHT + 0.02:
                problems.append(
                    f"slide {index}: a shape ends at {right:.2f} x {bottom:.2f}in"
                )

    used = set()
    for spec in SLIDES:
        for key in ("left", "right", "below", "body"):
            for block in spec.get(key, []):
                if block["type"] == "image":
                    used.add(Path(block["path"]).name)
                    if not (PROJECT_ROOT / block["path"]).exists():
                        problems.append(f"missing image {block['path']}")
    for figure in sorted((PROJECT_ROOT / "documents" / "deck-figures").glob("*.png")):
        if figure.name not in used:
            problems.append(f"unused figure {figure.name}")

    # The notes are a skeleton, not a script, so speaking them at a relaxed
    # 110 words per minute and allowing 12 seconds per slide for the pause,
    # the click and pointing at a chart gives the realistic floor.
    words = sum(len(notes_text(spec).split()) for spec in SLIDES)
    minutes = words / 110 + len(SLIDES) * 12 / 60
    counted = f"{len(presentation.slides)} slides"
    print(counted)  # noqa: T201
    print(f"speaker notes: {words} words, about {minutes:.0f} minutes delivered")  # noqa: T201
    if not 20 <= minutes <= 30:
        problems.append(
            f"estimated running time {minutes:.0f} minutes is outside the 20 to 30 minute target"
        )

    if problems:
        for problem in problems:
            print(f"  problem: {problem}")  # noqa: T201
        return 1
    print("all checks passed")  # noqa: T201
    return 0


if __name__ == "__main__":
    build()
    print(f"wrote {OUTPUT.relative_to(PROJECT_ROOT)}")  # noqa: T201
    if "--check" in sys.argv:
        raise SystemExit(check())
