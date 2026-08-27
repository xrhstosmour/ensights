#!/usr/bin/env python3
"""Build the defense cheatsheet from `defense_content.py`.

Usage:
    uv run python scripts/build-defense-cheatsheet.py

The cheatsheet walks the deck slide by slide, so it shares the content module
with `build-defense-deck.py`. Renumber or reorder a slide there and both
documents move together, which is the whole point: a walkthrough keyed to the
wrong slide numbers is worse than no walkthrough at all.
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.shared import RGBColor

sys.path.insert(0, str(Path(__file__).parent))

from defense_content import AUTHOR  # noqa: E402
from defense_content import BACKUP_SLIDES  # noqa: E402
from defense_content import KEY_NUMBERS  # noqa: E402
from defense_content import SECTIONS  # noqa: E402
from defense_content import SLIDES  # noqa: E402

PROJECT_ROOT = Path(__file__).parent.parent
OUTPUT = PROJECT_ROOT / "documents" / "ensights-defense-cheatsheet.docx"

GREEN = RGBColor(0x2C, 0x55, 0x30)
AMBER = RGBColor(0xC1, 0x82, 0x2A)
INK = RGBColor(0x24, 0x1F, 0x1B)
MUTED = RGBColor(0x6B, 0x62, 0x59)

HEADING_FONT = "Cambria"
BODY_FONT = "Calibri"

# Which slide index each section starts at, so the walkthrough can print the
# same four part headings the deck shows.
SECTION_STARTS = {
    index: spec["section"]
    for index, spec in enumerate(SLIDES)
    if spec["layout"] == "divider"
}


def paragraph(
    document,
    text,
    size=11,
    colour=INK,
    font=BODY_FONT,
    bold=False,
    before=0,
    after=6,
    align=None,
):
    block = document.add_paragraph()
    block.paragraph_format.space_before = Pt(before)
    block.paragraph_format.space_after = Pt(after)
    if align is not None:
        block.alignment = align
    run = block.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = colour
    return block


def slide_title(spec):
    if spec["layout"] == "divider":
        return f"Ενότητα: {SECTIONS[spec['section']][1]}"
    if spec["layout"] == "closing":
        return "Επίλογος και ερωτήσεις"
    return spec.get("title", "Τίτλος")


def add_walkthrough(document, entries, heading):
    paragraph(
        document,
        heading,
        size=15,
        colour=GREEN,
        font=HEADING_FONT,
        bold=True,
        before=20,
        after=10,
    )
    for number, spec in entries:
        paragraph(
            document,
            f"Διαφάνεια {number}: {slide_title(spec)}",
            size=13,
            colour=INK,
            font=HEADING_FONT,
            bold=True,
            before=14,
            after=4,
        )
        notes = spec["notes"]
        paragraph(document, "ΜΕ ΑΠΛΑ ΛΟΓΙΑ", size=9, colour=AMBER, bold=True, after=2)
        paragraph(document, notes["plain"])
        if notes.get("maths"):
            paragraph(
                document,
                "ΤΑ ΜΑΘΗΜΑΤΙΚΑ, Ή ΠΩΣ ΔΟΥΛΕΥΕΙ",
                size=9,
                colour=AMBER,
                bold=True,
                before=6,
                after=2,
            )
            paragraph(document, notes["maths"])
        if notes.get("question"):
            paragraph(
                document,
                "ΠΙΘΑΝΗ ΕΡΩΤΗΣΗ",
                size=9,
                colour=AMBER,
                bold=True,
                before=6,
                after=2,
            )
            paragraph(document, notes["question"])


def build():
    document = Document()
    style = document.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)

    paragraph(
        document,
        "ENSIGHTS · ΥΠΟΣΤΗΡΙΚΤΙΚΟ ΥΛΙΚΟ ΥΠΕΡΑΣΠΙΣΗΣ",
        size=10,
        colour=AMBER,
        bold=True,
        after=4,
    )
    paragraph(
        document,
        "Οδηγός διαφάνεια προς διαφάνεια",
        size=22,
        colour=GREEN,
        font=HEADING_FONT,
        bold=True,
        after=6,
    )
    paragraph(
        document,
        "Κάθε διαφάνεια της παρουσίασης, εξηγημένη απλά, με τα μαθηματικά "
        "όπου υπάρχουν και τις ερωτήσεις που μπορεί να κάνει η επιτροπή.",
        size=12,
        colour=MUTED,
        after=10,
    )
    paragraph(document, f"{AUTHOR} · ΕΚΠΑ", size=10, colour=MUTED, after=14)
    paragraph(
        document,
        "Πώς να το χρησιμοποιήσεις: κάθε διαφάνεια έχει το δικό της κομμάτι "
        "εδώ, με την ίδια σειρά και την ίδια αρίθμηση. Πρώτα μια εξήγηση σαν "
        "να τη λες σε κάποιον που δεν ξέρει τίποτα από ανάλυση δεδομένων, "
        "μετά τα μαθηματικά αν υπάρχουν, και μετά η ερώτηση που μπορεί να "
        "σου γίνει πάνω σε αυτή τη διαφάνεια. Στο τέλος υπάρχει κάρτα με όλα "
        "τα νούμερα, για γρήγορη ανάκληση.",
    )

    # The title and contents slides come before the first divider, so the
    # walkthrough opens with its own heading and then follows the deck's parts.
    heading = "Εισαγωγή"
    pending = []
    for index, spec in enumerate(SLIDES):
        if index in SECTION_STARTS:
            if pending:
                add_walkthrough(document, pending, heading)
                pending = []
            marker, name, _ = SECTIONS[SECTION_STARTS[index]]
            heading = f"Μέρος {marker}: {name}"
        pending.append((index + 1, spec))
    if pending:
        add_walkthrough(document, pending, heading)

    add_walkthrough(
        document,
        [(f"B{index + 1}", spec) for index, spec in enumerate(BACKUP_SLIDES)],
        "Εφεδρικές διαφάνειες, μετά το ευχαριστώ",
    )

    document.add_page_break()
    paragraph(
        document,
        "Κάρτα νούμερων",
        size=15,
        colour=GREEN,
        font=HEADING_FONT,
        bold=True,
        after=4,
    )
    paragraph(
        document,
        "Για γρήγορη ανάκληση πριν μπεις, και για όταν η επιτροπή ρωτήσει έναν αριθμό.",
        size=11,
        colour=MUTED,
        after=10,
    )
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in KEY_NUMBERS:
        cells = table.add_row().cells
        for cell, text, bold in ((cells[0], label, True), (cells[1], value, False)):
            block = cell.paragraphs[0]
            block.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = block.add_run(text)
            run.font.size = Pt(10)
            run.font.bold = bold
            run.font.name = BODY_FONT

    document.save(OUTPUT)
    return document


def check():
    document = Document(OUTPUT)
    text = "\n".join(block.text for block in document.paragraphs)
    problems = []
    for index, spec in enumerate(SLIDES, start=1):
        if f"Διαφάνεια {index}: {slide_title(spec)}" not in text:
            problems.append(f"slide {index} missing from the walkthrough")
    for index in range(1, len(BACKUP_SLIDES) + 1):
        if f"Διαφάνεια B{index}:" not in text:
            problems.append(f"backup slide B{index} missing from the walkthrough")
    rows = sum(len(table.rows) for table in document.tables)
    if rows != len(KEY_NUMBERS):
        problems.append(f"numbers card has {rows} rows, expected {len(KEY_NUMBERS)}")

    print(f"{len(SLIDES)} slides plus {len(BACKUP_SLIDES)} backup, {rows} numbers")  # noqa: T201
    if problems:
        for problem in problems:
            print(f"  problem: {problem}")  # noqa: T201
        return 1
    print("all checks passed")  # noqa: T201
    return 0


if __name__ == "__main__":
    build()
    print(f"wrote {OUTPUT.relative_to(PROJECT_ROOT)}")  # noqa: T201
    if "--check" in sys.argv:
        raise SystemExit(check())
